"""本地 API：上传 PDF/PPTX → 按页提取（文字 + 版式/插图）→ DeepSeek 教学批注 → 预览编辑 → 导出。"""

from __future__ import annotations

import base64
import io
import os
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from openai import OpenAI
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.slide import Slide
from pypdf import PdfReader

app = FastAPI(title="SlideAnnotate Local API", version="0.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_ROOT = Path(tempfile.gettempdir()) / "slide-annotate-uploads"
UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)

_jobs: dict[str, dict[str, Any]] = {}

# 单图最长边像素（控制体积与 token）
_MAX_VISUAL_SIDE = int(os.environ.get("SLIDE_ANNOTATE_IMAGE_MAX_SIDE", "1280"))
# 界面预览图最长边（可高于模型用图，仅影响预览清晰度）
_PREVIEW_MAX_SIDE = int(os.environ.get("SLIDE_ANNOTATE_PREVIEW_MAX_SIDE", "2400"))
# 单页最多附加几张「内嵌图」（PPT 形状图；PDF 另附整页渲染）
_MAX_INLINE_IMAGES = int(os.environ.get("SLIDE_ANNOTATE_MAX_INLINE_IMAGES", "4"))


def _job_path(job_id: str) -> Path:
    return UPLOAD_ROOT / job_id


def _extract_pdf_text(path: Path) -> list[str]:
    reader = PdfReader(str(path))
    out: list[str] = []
    for page in reader.pages:
        t = page.extract_text() or ""
        out.append(t.strip())
    return out


def _slide_visible_text(slide: Slide) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs)
            if line.strip():
                parts.append(line.strip())
    return "\n".join(parts).strip()


def _extract_pptx_text(path: Path) -> list[str]:
    prs = Presentation(str(path))
    return [_slide_visible_text(s) for s in prs.slides]


def _pptx_slide_inline_images(slide: Slide) -> list[bytes]:
    blobs: list[bytes] = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            blobs.append(shape.image.blob)
        except Exception:
            continue
        if len(blobs) >= _MAX_INLINE_IMAGES:
            break
    return blobs


def _build_pdf_page_renders(path: Path, n_pages: int) -> list[bytes | None]:
    """每页渲染为 PNG（版式、图表、扫描件均可被模型看到）。"""
    return _render_pdf_pages_png(path, n_pages, _MAX_VISUAL_SIDE)


def _render_pdf_pages_png(path: Path, n_pages: int, max_side: int) -> list[bytes | None]:
    out: list[bytes | None] = []
    doc = fitz.open(str(path))
    try:
        count = min(n_pages, doc.page_count)
        for i in range(count):
            out.append(_render_pdf_page_png(doc, i, max_side))
        while len(out) < n_pages:
            out.append(None)
    finally:
        doc.close()
    return out


def _render_pdf_page_png(doc: fitz.Document, page_index: int, max_side: int) -> bytes | None:
    if page_index < 0 or page_index >= doc.page_count:
        return None
    page = doc[page_index]
    w, h = page.rect.width, page.rect.height
    if w < 1 or h < 1:
        return None
    scale = min(max_side / max(w, h), 3.0)
    mat = fitz.Matrix(scale, scale)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    return pix.tobytes("png")


def _build_pptx_inline_images(path: Path) -> list[list[bytes]]:
    prs = Presentation(str(path))
    return [_pptx_slide_inline_images(s) for s in prs.slides]


def _soffice_candidates() -> list[str]:
    env = os.environ.get("SLIDE_ANNOTATE_SOFFICE", "").strip()
    names = ["soffice", "libreoffice"]
    out: list[str] = []
    if env:
        out.append(env)
    for name in names:
        found = shutil.which(name)
        if found:
            out.append(found)
    if os.name == "nt":
        for base in (
            os.environ.get("ProgramFiles", r"C:\Program Files"),
            os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
        ):
            exe = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if exe.is_file():
                out.append(str(exe))
    seen: set[str] = set()
    deduped: list[str] = []
    for p in out:
        if p not in seen:
            seen.add(p)
            deduped.append(p)
    return deduped


def _try_pptx_to_pdf(pptx_path: Path, out_dir: Path) -> Path | None:
    """PPTX → PDF（LibreOffice），供整页幻灯片预览。未安装 LO 时返回 None。"""
    for soffice in _soffice_candidates():
        try:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--nologo",
                    "--nofirststartwizard",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(out_dir),
                    str(pptx_path),
                ],
                check=True,
                timeout=180,
                capture_output=True,
            )
            pdf_path = out_dir / f"{pptx_path.stem}.pdf"
            if pdf_path.is_file():
                return pdf_path
        except Exception:
            continue
    return None


def _guess_image_media(raw: bytes) -> str:
    if raw[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if raw[:2] == b"\xff\xd8":
        return "image/jpeg"
    if raw[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    return "image/png"


def _downscale_image_bytes(raw: bytes, max_side: int = _MAX_VISUAL_SIDE) -> bytes:
    """统一压缩为 JPEG，减小多模态 payload。"""
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(raw))
        im = im.convert("RGB")
        w, h = im.size
        if max(w, h) > max_side:
            im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=85, optimize=True)
        return buf.getvalue()
    except Exception:
        return raw


def _bytes_to_data_url(raw: bytes) -> str:
    if len(raw) >= 2 and raw[:2] == b"\xff\xd8":
        mime = "image/jpeg"
        payload = raw
    elif raw[:8] == b"\x89PNG\r\n\x1a\n":
        mime = "image/png"
        payload = raw
    else:
        mime = "image/jpeg"
        payload = _downscale_image_bytes(raw)
    b64 = base64.standard_b64encode(payload).decode("ascii")
    return f"data:{mime};base64,{b64}"


def _visuals_to_data_urls(visual: dict[str, Any]) -> list[str]:
    urls: list[str] = []
    primary = visual.get("primary")
    if primary:
        urls.append(_bytes_to_data_url(_downscale_image_bytes(primary)))
    for blob in visual.get("inline") or []:
        if not blob:
            continue
        urls.append(_bytes_to_data_url(blob))
        if len(urls) >= 1 + _MAX_INLINE_IMAGES:
            break
    return urls[: 1 + _MAX_INLINE_IMAGES]


SYSTEM_PROMPT = """你是面对中国学生的讲师助手。用户会给你某一页英文讲义（来自 PDF 或 PPT）的正文摘录，并可能附上：
- PDF：整页版式截图（含图表、公式、流程图、扫描文字）；
- PPT：幻灯片上的插图/截图。

请结合**文字与图像**用中文写出「教学批注」：帮助学生理解这一页在讲什么、关键概念、图表/图示含义、常见误区、课堂提问建议、与前后内容的衔接。
要求：
- 不要整页逐句翻译英文；允许必要时引用原文术语并马上用中文解释。
- 若图中有坐标轴、图例、关键数字，请在批注中点明其作用，但不要编造图中没有的数值。
- 语气简洁、可当讲师口播提纲使用，300–700 字为宜（信息少的页面可以更短）。
- 不要编造页面上没有出现过的具体数据、人名、公式或结论；不确定就写「需结合全文确认」。"""


def _call_deepseek(
    api_key: str,
    user_text: str,
    image_data_urls: list[str] | None = None,
) -> str:
    base = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com").rstrip("/")
    if image_data_urls:
        model = os.environ.get("DEEPSEEK_VISION_MODEL") or os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
    else:
        model = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
    client = OpenAI(api_key=api_key, base_url=base)

    user_content: str | list[dict[str, Any]]
    if image_data_urls:
        user_content = [{"type": "text", "text": user_text}]
        for url in image_data_urls:
            user_content.append({"type": "image_url", "image_url": {"url": url}})
    else:
        user_content = user_text

    def _do_call(m: str) -> str:
        r = client.chat.completions.create(
            model=m,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_content},
            ],
            temperature=0.4,
        )
        msg = r.choices[0].message
        if not msg or not msg.content:
            return ""
        return msg.content.strip()

    try:
        return _do_call(model)
    except Exception as e:
        err = str(e).lower()
        # 当前模型不支持多模态或图片过大时，回退纯文本
        if image_data_urls and (
            "image" in err
            or "multimodal" in err
            or "vision" in err
            or "content" in err
            or "400" in err
            or "not support" in err
            or "unsupported" in err
        ):
            fallback_model = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
            user_content = (
                user_text
                + "\n\n（注意：接口未接受附图或当前模型不支持读图，请仅根据上述文字做批注，并提醒学生对照原稿中的图表。）"
            )
            r2 = client.chat.completions.create(
                model=fallback_model,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_content},
                ],
                temperature=0.4,
            )
            msg = r2.choices[0].message
            if not msg or not msg.content:
                return ""
            return msg.content.strip()
        raise


def _ensure_job(job_id: str) -> dict[str, Any]:
    if job_id not in _jobs:
        raise HTTPException(status_code=404, detail="任务不存在或已过期")
    return _jobs[job_id]


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/jobs")
async def create_job(file: UploadFile = File(...)) -> JSONResponse:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".pdf", ".pptx"):
        raise HTTPException(
            status_code=400,
            detail="仅支持 .pdf 与 .pptx（请将 .ppt 另存为 .pptx）",
        )

    job_id = str(uuid.uuid4())
    job_dir = _job_path(job_id)
    job_dir.mkdir(parents=True, exist_ok=False)

    raw_name = f"original{suffix}"
    raw_path = job_dir / raw_name
    try:
        with raw_path.open("wb") as f:
            shutil.copyfileobj(file.file, f)
    finally:
        await file.close()

    if suffix == ".pdf":
        texts = _extract_pdf_text(raw_path)
        doc_type = "pdf"
        renders = _build_pdf_page_renders(raw_path, len(texts))
        visuals: list[dict[str, Any]] = [{"primary": r, "inline": []} for r in renders]
    else:
        texts = _extract_pptx_text(raw_path)
        doc_type = "pptx"
        inline_lists = _build_pptx_inline_images(raw_path)
        pdf_proxy = _try_pptx_to_pdf(raw_path, job_dir)
        if pdf_proxy:
            renders = _build_pdf_page_renders(pdf_proxy, len(texts))
            visuals = [
                {"primary": renders[i] if i < len(renders) else None, "inline": inline_lists[i]}
                for i in range(len(texts))
            ]
        else:
            visuals = [{"primary": None, "inline": inline} for inline in inline_lists]

    pages: list[dict[str, Any]] = []
    for i, t in enumerate(texts):
        v = visuals[i] if i < len(visuals) else {"primary": None, "inline": []}
        note_parts: list[str] = []
        if v.get("primary"):
            note_parts.append("含整页版式图（图表/版式已提交模型）")
        n_in = len(v.get("inline") or [])
        if n_in:
            note_parts.append(f"含 {n_in} 张幻灯片内嵌图（已提交模型）")
        image_note = "；".join(note_parts)
        pages.append(
            {
                "index": i,
                "source_text": t,
                "image_note": image_note,
                "annotation": "",
                "has_preview": bool(v.get("primary") or v.get("inline")),
            }
        )

    _jobs[job_id] = {
        "id": job_id,
        "type": doc_type,
        "original_path": str(raw_path),
        "original_name": file.filename or raw_name,
        "pages": pages,
        "_visuals": visuals,
    }

    return JSONResponse(
        {
            "job_id": job_id,
            "type": doc_type,
            "page_count": len(pages),
            "pages": pages,
        }
    )


@app.post("/api/jobs/{job_id}/generate")
async def generate_annotations(job_id: str, payload: dict[str, Any]) -> JSONResponse:
    job = _ensure_job(job_id)
    api_key = (payload or {}).get("api_key") or os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        raise HTTPException(status_code=400, detail="缺少 DeepSeek API Key（请在界面填写或设置环境变量 DEEPSEEK_API_KEY）")

    indices = (payload or {}).get("indices")
    pages: list[dict[str, Any]] = job["pages"]
    visuals: list[dict[str, Any]] = job.get("_visuals") or []

    if indices is None:
        targets = list(range(len(pages)))
    else:
        targets = [int(i) for i in indices if 0 <= int(i) < len(pages)]

    for i in targets:
        p = pages[i]
        src = p["source_text"] or ""
        img_hint = p.get("image_note") or ""
        meta_lines = [f"这是第 {i + 1} 页的正文摘录：", "", src or "（本页无可用文本层，请主要依据附图理解。）"]
        if img_hint:
            meta_lines.extend(["", f"图像说明：{img_hint}"])
        user_text = "\n".join(meta_lines)

        vis = visuals[i] if i < len(visuals) else {}
        data_urls = _visuals_to_data_urls(vis) if vis else []

        try:
            pages[i]["annotation"] = _call_deepseek(str(api_key), user_text, data_urls or None)
        except Exception as e:
            pages[i]["annotation"] = f"[生成失败] {e}"

    job["pages"] = pages
    return JSONResponse({"pages": pages})


@app.put("/api/jobs/{job_id}/pages/{page_index}")
async def update_page(job_id: str, page_index: int, payload: dict[str, Any]) -> JSONResponse:
    job = _ensure_job(job_id)
    pages: list[dict[str, Any]] = job["pages"]
    if page_index < 0 or page_index >= len(pages):
        raise HTTPException(status_code=400, detail="页码无效")
    ann = (payload or {}).get("annotation")
    if ann is not None:
        pages[page_index]["annotation"] = str(ann)
    return JSONResponse({"page": pages[page_index]})


@app.api_route("/api/jobs/{job_id}/export", methods=["GET", "POST"])
def export_job(job_id: str) -> FileResponse:
    job = _ensure_job(job_id)
    job_dir = _job_path(job_id)
    annotations = [p.get("annotation") or "" for p in job["pages"]]
    src = Path(job["original_path"])

    if job["type"] == "pdf":
        out = job_dir / "annotated.pdf"
        _export_pdf_with_notes(src, annotations, out)
        media = "application/pdf"
        filename = f"annotated_{Path(job['original_name']).stem}.pdf"
    else:
        out = job_dir / "annotated.pptx"
        _export_pptx_notes(src, annotations, out)
        media = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        filename = f"annotated_{Path(job['original_name']).stem}.pptx"

    return FileResponse(
        path=str(out),
        media_type=media,
        filename=filename,
    )


def _render_note_page(page: fitz.Page, page_num: int, text: str) -> None:
    """独立「教学批注」页：纯页面内容，不依赖 PDF 注释/弹窗（兼容 Edge/Chrome）。"""
    r = page.rect
    page.draw_rect(r, fill=(0.97, 0.94, 0.99), color=(0.82, 0.72, 0.92), width=1.0)

    title_rect = fitz.Rect(48, 42, r.width - 48, 88)
    body_rect = fitz.Rect(48, 96, r.width - 48, r.height - 48)
    fonts = ("china-s", "china-ss", "helv")

    for fname in fonts:
        try:
            page.insert_textbox(
                title_rect,
                f"第 {page_num} 页 · 教学批注",
                fontsize=14,
                fontname=fname,
                color=(0.35, 0.2, 0.55),
                align=fitz.TEXT_ALIGN_LEFT,
            )
            break
        except Exception:
            continue

    base_fs = 10.0 if len(text) > 600 else 11.0
    for fname in fonts:
        fs = base_fs
        while fs >= 7.5:
            try:
                rv = page.insert_textbox(
                    body_rect,
                    text,
                    fontsize=fs,
                    fontname=fname,
                    color=(0.15, 0.13, 0.1),
                    align=fitz.TEXT_ALIGN_LEFT,
                )
                if rv >= 0:
                    return
            except Exception:
                break
            fs -= 0.5

    page.insert_text((52, 110), text[:4000], fontsize=8, fontname="helv", color=(0.15, 0.13, 0.1))


def _export_pdf_with_notes(src: Path, notes: list[str], out: Path) -> None:
    """导出 PDF：每页原文后插入一页「教学批注」说明（若该页有批注）。

    不使用 Text/Popup 注释层——Edge/Chrome 内置查看器对注释悬停/点击支持差，
    改为普通页面内容，任意阅读器均可稳定阅读。
    """
    src_doc = fitz.open(str(src))
    out_doc = fitz.open()
    try:
        n = src_doc.page_count
        for i in range(n):
            out_doc.insert_pdf(src_doc, from_page=i, to_page=i)
            text = (notes[i] if i < len(notes) else "").strip()
            if not text:
                continue
            src_page = src_doc[i]
            w, h = src_page.rect.width, src_page.rect.height
            note_page = out_doc.new_page(width=w, height=h)
            _render_note_page(note_page, i + 1, text)
        out_doc.save(str(out), garbage=4, deflate=True)
    finally:
        src_doc.close()
        out_doc.close()


def _export_pptx_notes(src: Path, notes: list[str], out: Path) -> None:
    prs = Presentation(str(src))
    for i, slide in enumerate(prs.slides):
        note = notes[i] if i < len(notes) else ""
        ns = slide.notes_slide
        ns.notes_text_frame.text = note
    prs.save(str(out))


@app.get("/api/jobs/{job_id}/pages/{page_index}/preview")
def page_preview(job_id: str, page_index: int) -> Response:
    """返回单页预览图（高分辨率）：PDF/PPTX(经 LO 转 PDF) 实时渲染；否则回退内嵌图。"""
    job = _ensure_job(job_id)
    pages: list[dict[str, Any]] = job["pages"]
    if page_index < 0 or page_index >= len(pages):
        raise HTTPException(status_code=404, detail="页不存在")

    if job["type"] == "pdf":
        png = _render_pdf_page_png_from_path(Path(job["original_path"]), page_index, _PREVIEW_MAX_SIDE)
        if png:
            return Response(content=png, media_type="image/png")

    job_dir = _job_path(job_id)
    pdf_proxy = job_dir / f"{Path(job['original_path']).stem}.pdf"
    if pdf_proxy.is_file():
        png = _render_pdf_page_png_from_path(pdf_proxy, page_index, _PREVIEW_MAX_SIDE)
        if png:
            return Response(content=png, media_type="image/png")

    visuals: list[dict[str, Any]] = job.get("_visuals") or []
    if page_index < len(visuals):
        vis = visuals[page_index]
        primary = vis.get("primary")
        if primary:
            return Response(content=primary, media_type="image/png")
        inline: list[bytes] = vis.get("inline") or []
        if inline:
            raw = inline[0]
            return Response(content=raw, media_type=_guess_image_media(raw))

    raise HTTPException(status_code=404, detail="该页暂无预览图（PPTX 可安装 LibreOffice 后重新上传以生成整页预览）")


def _render_pdf_page_png_from_path(path: Path, page_index: int, max_side: int) -> bytes | None:
    if not path.is_file():
        return None
    doc = fitz.open(str(path))
    try:
        return _render_pdf_page_png(doc, page_index, max_side)
    finally:
        doc.close()


@app.get("/api/jobs/{job_id}/meta")
def job_meta(job_id: str) -> dict[str, Any]:
    job = _ensure_job(job_id)
    return {
        "job_id": job["id"],
        "type": job["type"],
        "original_name": job["original_name"],
        "page_count": len(job["pages"]),
    }


@app.get("/api/jobs/{job_id}/pages")
def job_pages(job_id: str) -> dict[str, Any]:
    job = _ensure_job(job_id)
    return {"pages": job["pages"]}


def main() -> None:
    import uvicorn

    host = os.environ.get("SLIDE_ANNOTATE_HOST", "127.0.0.1")
    port = int(os.environ.get("SLIDE_ANNOTATE_PORT", "8765"))
    uvicorn.run("app:app", host=host, port=port, reload=False)


if __name__ == "__main__":
    main()
